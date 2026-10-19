# -*- coding: utf-8 -*-
"""בונה דו"ח הדגמה עם מטופל בדיוני ותמונות סינתטיות - בלי לוגו, בלי פרטי מרפאה ובלי נתוני מטופלים אמיתיים."""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

IMG = sys.argv[1]
OUT = sys.argv[2]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:cs"), "Arial")
st.element.rPr.get_or_add_szCs() if hasattr(st.element.rPr, "get_or_add_szCs") else None


def rtl(p):
    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    for r in p.runs:
        r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    return p


def para(text, bold=False, center=False, underline=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.underline = underline
    if size:
        r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return rtl(p)


def table(rows, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rtl(p)
            if i % 2 == 0 and val:
                for r in p.runs:
                    r.bold = True
    return t


para("קבלה משולבת - דו\"ח הדגמה", bold=True, center=True, size=16)
para("כל הנתונים בדו\"ח זה בדיוניים ונוצרו לצורך הדגמה בלבד", center=True, size=9)
table([["שם", "ת.ז", "ת לידה", "טל", "מייל"],
       ["ישראל ישראלי (בדיוני)", "000000000", "01/01/1980", "050-0000000", "demo@example.com"]])
para("")
para("אבחנות", bold=True)
para("NSLBP\nPostural Imbalance\nPelvic Weakness")
para("כללי", bold=True)
para("בן 46, עובד משרד, כאבי גב תחתון מזה שנתיים, מוחמרים בעמידה ממושכת ובסוף יום עבודה. "
     "ללא אירוע טראומטי. פעילות גופנית: הליכות פעמיים בשבוע. ללא תרופות קבועות.")
para("בדיקה פיזיקלית", bold=True)
para("הגבלה קלה בכיפוף ירך ימין (110° לעומת 120°), קיצור דו-צדדי של ה-Hamstrings (SLR 60°), "
     "רגישות קלה מעל SIJ ימין. כוח ותחושה שמורים.")
para("")
para("מדדים סטטים בעמידה :", bold=True, center=True, underline=True)
table([["Kyphosis", "Scoliosis", "Lordosis"], ["", "", ""],
       ["Coronal Imbalance", "Sagittal  Imbalance", "Fleche Cervicale"],
       ['הסטה של האגן  מ"מ', 'מ"מ הסטה של הגוף  קדימה', 'צווארי  מ"מ\nמותני  מ"מ'],
       ["Pelvic Rotation", "Pelvic Torsion", "Pelvic Obliquity"], ["מעלות", "נורמה", "הרמה אגנית"]])
para("")
para("Coronal view", bold=True, center=True, underline=True)
doc.add_picture(os.path.join(IMG, "demo_diers_coronal.png"), width=Cm(16.5))
para("")
para("בדיקת Stabilimeter  עניים פתוחות ועצומות :", bold=True, center=True, underline=True)
para("1. בדיקת עמידה וחלוקת עומסים  ב   Sagital & Coronal View", bold=True, underline=True)
para("עיניים פקוחות (ירוק) / עיניים עצומות (אדום)", size=9)
doc.add_picture(os.path.join(IMG, "demo_pedoscan.png"), width=Cm(16.5))
para("")
para("דיון :", bold=True, underline=True)
para("המלצות:", bold=True, underline=True)
para("פיזיותרפיה - חיזוק אגן ועבודה על טווחים.")
doc.save(OUT)
print("saved", OUT)
